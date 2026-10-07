# -*- coding: utf-8 -*-
"""
Script tạo 2 tài liệu chuẩn cho dự án Web Quản lý Hợp đồng Bảo hiểm:
1. SRS_Insurance_Policy_Management_System.docx (Tài liệu Đặc tả Yêu cầu Phần mềm)
2. FDD_System_Workflow_and_Feature_Specification.docx (Tài liệu Đặc tả Luồng Hoạt động & Danh mục Chức năng Hệ thống)
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# ==========================================================
# COLOR PALETTE (Chuyên nghiệp, Tinh tế, Hiện đại)
# ==========================================================
COLOR_PRIMARY_DARK = RGBColor(26, 54, 93)     # Deep Navy #1A365D
COLOR_PRIMARY_BLUE = RGBColor(37, 99, 235)    # Brand Blue #2563EB
COLOR_TEAL_ACCENT  = RGBColor(13, 148, 136)   # Teal #0D9488
COLOR_TEXT_MAIN    = RGBColor(30, 41, 59)     # Slate 800 #1E293B
COLOR_TEXT_MUTED   = RGBColor(100, 116, 139)  # Slate 500 #64748B
COLOR_WARNING      = RGBColor(217, 119, 6)    # Amber 600 #D97706
COLOR_SUCCESS      = RGBColor(22, 163, 74)    # Green 600 #16A34A

HEX_PRIMARY_DARK = "1A365D"
HEX_PRIMARY_BLUE = "2563EB"
HEX_BG_LIGHT     = "F8FAFC"
HEX_BORDER_LIGHT = "CBD5E1"
HEX_HEADER_BG    = "1E3A8A"
HEX_ACCENT_BG    = "EFF6FF"
HEX_CALLOUT_BG   = "F1F5F9"
HEX_WARNING      = "D97706"
HEX_SUCCESS      = "16A34A"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def format_cell_borders(cell, border_color="CBD5E1"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:left w:val="none"/><w:right w:val="none"/></w:tcBorders>')
    tcPr.append(borders)

def add_styled_title(doc, title_text, subtitle_text=None, meta_info=None):
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(title_text)
    run.font.name = 'Calibri'
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY_DARK

    if subtitle_text:
        sub_p = doc.add_paragraph()
        sub_p.paragraph_format.space_before = Pt(0)
        sub_p.paragraph_format.space_after = Pt(16)
        sub_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub_p.add_run(subtitle_text)
        sub_run.font.name = 'Calibri'
        sub_run.font.size = Pt(13)
        sub_run.font.italic = True
        sub_run.font.color.rgb = COLOR_TEXT_MUTED

    if meta_info:
        card_table = doc.add_table(rows=1, cols=1)
        card_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = card_table.rows[0].cells[0]
        set_cell_background(cell, HEX_ACCENT_BG)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:bottom w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_PRIMARY_BLUE}"/><w:right w:val="none"/></w:tcBorders>')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        for line in meta_info:
            r = p.add_run(line + "\n")
            r.font.name = 'Calibri'
            r.font.size = Pt(10)
            r.font.color.rgb = COLOR_TEXT_MAIN
        # Remove trailing newline from last run
        if p.runs:
            p.runs[-1].text = p.runs[-1].text.rstrip()

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY_DARK
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY_BLUE
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_TEAL_ACCENT
    return p

def add_body_paragraph(doc, text, bold_prefix=None, italic=False, space_after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.font.name = 'Calibri'
        r_prefix.font.size = Pt(10.5)
        r_prefix.font.bold = True
        r_prefix.font.color.rgb = COLOR_TEXT_MAIN
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10.5)
    r.font.italic = italic
    r.font.color.rgb = COLOR_TEXT_MAIN
    return p

def add_bullet_item(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.font.name = 'Calibri'
        r_prefix.font.size = Pt(10.5)
        r_prefix.font.bold = True
        r_prefix.font.color.rgb = COLOR_TEXT_MAIN
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10.5)
    r.font.color.rgb = COLOR_TEXT_MAIN
    return p

def add_callout(doc, text, title="LƯU Ý NGHIỆP VỤ QUAN TRỌNG:", border_color=HEX_PRIMARY_BLUE, bg_color=HEX_CALLOUT_BG):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.rows[0].cells[0]
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:bottom w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:right w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    if title:
        r_title = p.add_run(title + " ")
        r_title.font.name = 'Calibri'
        r_title.font.size = Pt(10)
        r_title.font.bold = True
        r_title.font.color.rgb = COLOR_PRIMARY_DARK
        
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_TEXT_MAIN
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_table_data(doc, headers, data_rows, col_widths=None):
    table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        cell = hdr_cells[i]
        set_cell_background(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=140, bottom=140, left=120, right=120)
        format_cell_borders(cell, HEX_HEADER_BG)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(header_text)
        r.font.name = 'Calibri'
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    # Data Rows
    for row_idx, row_data in enumerate(data_rows):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = HEX_BG_LIGHT if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            cell = row_cells[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            format_cell_borders(cell, HEX_BORDER_LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(cell_value))
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)
            r.font.color.rgb = COLOR_TEXT_MAIN
            
    # Apply Column Widths if provided
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Inches(width)
                    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table

print("Setup helper functions completed!")
